import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- COLORS ---
NAVY_PRIMARY = RGBColor(30, 39, 97)      # #1E2761
NAVY_SECONDARY = RGBColor(20, 27, 77)    # #141B4D
CORAL = RGBColor(255, 107, 91)           # #FF6B5B
GOLD = RGBColor(249, 231, 149)           # #F9E795
MINT = RGBColor(47, 208, 138)            # #2FD08A
ICE = RGBColor(202, 220, 252)            # #CADCFC
WHITE = RGBColor(255, 255, 255)          # #FFFFFF
RED_TINT = RGBColor(60, 20, 40)          # for comparison table

# --- SETUP ---
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# --- HELPERS ---
def add_bg(slide, color=NAVY_PRIMARY):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_text(slide, text, left, top, width, height, font_size, color=WHITE, font_name='Calibri', bold=False, italic=False, align=PP_ALIGN.LEFT, word_wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.name = font_name
    p.font.bold = bold
    p.font.italic = italic
    p.alignment = align
    return txBox

def add_card(slide, left, top, width, height, color=NAVY_SECONDARY, border_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
    else:
        card.line.fill.background()
    return card

def add_circle_icon(slide, left, top, size, char, bg_color=NAVY_PRIMARY):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = bg_color
    circle.line.color.rgb = WHITE
    circle.line.width = Pt(1)
    
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = char
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    return circle

def add_header(slide, kicker, headline, subhead):
    add_text(slide, kicker.upper(), Inches(0.8), Inches(0.6), Inches(10), Inches(0.4), 12, CORAL, bold=True)
    add_text(slide, headline, Inches(0.8), Inches(0.9), Inches(11.5), Inches(0.8), 36, WHITE, 'Cambria', bold=True)
    if subhead:
        add_text(slide, subhead, Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.6), 20, ICE, italic=True)

def add_page_number(slide, num):
    add_text(slide, f"{num} / 9", Inches(12.2), Inches(6.9), Inches(1), Inches(0.4), 12, ICE, align=PP_ALIGN.RIGHT)


# ==========================================
# SLIDE 1: COVER
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
add_bg(s1)
# Decor circles
add_circle_icon(s1, Inches(10), Inches(-1), Inches(4), "", NAVY_SECONDARY)
add_circle_icon(s1, Inches(11), Inches(5), Inches(3), "", NAVY_SECONDARY)

add_text(s1, "M→M", Inches(0.8), Inches(0.8), Inches(2), Inches(0.6), 24, WHITE, 'Cambria', bold=True)
add_text(s1, "PROBLEM VERTICAL · AGENTIC AI · CODE KUDLA 2026", Inches(0.8), Inches(1.2), Inches(6), Inches(0.4), 11, CORAL, bold=True)

add_text(s1, "MeetingToMotion", Inches(0.8), Inches(2.2), Inches(10), Inches(1), 54, WHITE, 'Cambria', bold=True)
add_text(s1, "Autonomous Cross-Tool Action-Item Executor", Inches(0.8), Inches(3.2), Inches(10), Inches(0.6), 28, ICE)
add_text(s1, "Turns meeting decisions into completed work across Jira, Gmail and Notion — and knows when to stop and ask a human instead of guessing.", Inches(0.8), Inches(3.8), Inches(7.5), Inches(1), 18, ICE)

add_card(s1, Inches(0.8), Inches(5.2), Inches(6.5), Inches(1.6))
add_text(s1, "TEAM DETAILS", Inches(1.0), Inches(5.4), Inches(6), Inches(0.3), 12, CORAL, bold=True)
add_text(s1, "Team Name: Sea Of Syntax", Inches(1.0), Inches(5.8), Inches(6), Inches(0.4), 16, WHITE, bold=True)
add_text(s1, "Members: Aadish Balakrishna Salian, Amish Sudhakara, Hardik Shetty, Aadithya Deepak M", Inches(1.0), Inches(6.2), Inches(6), Inches(0.4), 14, ICE)

add_card(s1, Inches(7.8), Inches(5.2), Inches(4.5), Inches(1.6))
add_text(s1, "\"Extraction is commoditized.\nExecution is the moat.\"", Inches(8.0), Inches(5.6), Inches(4.1), Inches(1.0), 22, GOLD, 'Cambria', italic=True, align=PP_ALIGN.CENTER)
add_page_number(s1, 1)

# ==========================================
# SLIDE 2: THE PROBLEM
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
add_bg(s2)
add_header(s2, "The Problem", "Meeting bots stop exactly where the work begins.", "A perfect summary still leaves your team with a slow, manual translation job.")

# Left Card
add_card(s2, Inches(0.8), Inches(2.4), Inches(5.5), Inches(4.2))
add_text(s2, "A Monday morning, everywhere", Inches(1.1), Inches(2.7), Inches(5), Inches(0.5), 20, WHITE, bold=True)
add_text(s2, "The standup ends. The AI notetaker drops a tidy summary in Slack. Four things need to happen — a ticket, an email, a doc update, one unclear owner. And then... someone still has to actually go do all of it, one tab at a time, later today, if they remember.", Inches(1.1), Inches(3.2), Inches(5), Inches(1.5), 16, ICE)
add_text(s2, "Sound familiar? That gap between 'we decided' and 'it's done' is where momentum quietly dies.", Inches(1.1), Inches(4.8), Inches(5), Inches(0.8), 16, GOLD, italic=True)
add_text(s2, "Existing tools (Otter, Fireflies, Fathom, Zoom AI) already nail extraction. That part is commoditized.", Inches(1.1), Inches(5.6), Inches(5), Inches(0.6), 14, ICE)

# Right list
add_text(s2, "THE TRANSLATION GAP", Inches(6.8), Inches(2.4), Inches(5), Inches(0.4), 14, CORAL, bold=True)

y_start = 2.9
gap = 0.8
items = [
    ("T", "Create the Jira ticket", "Pick project, assignee, priority, write the description by hand"),
    ("M", "Draft the follow-up email", "Translate discussion into a recipient-ready message"),
    ("D", "Update the project doc", "Move decisions and references into Notion manually"),
    ("?", "Chase down the ambiguity", "Find out which 'Sarah' or which deadline they actually meant")
]

for i, (icon, title, desc) in enumerate(items):
    add_circle_icon(s2, Inches(6.8), Inches(y_start + (i*gap)), Inches(0.5), icon)
    add_text(s2, title, Inches(7.5), Inches(y_start + (i*gap)), Inches(5), Inches(0.3), 16, WHITE, bold=True)
    add_text(s2, desc, Inches(7.5), Inches(y_start + (i*gap) + 0.3), Inches(5), Inches(0.4), 14, ICE)

add_card(s2, Inches(6.8), Inches(6.1), Inches(5.7), Inches(0.6), NAVY_SECONDARY, CORAL)
add_text(s2, "Result: slower execution, missed commitments, zero visibility into what actually got done.", Inches(7.0), Inches(6.25), Inches(5.3), Inches(0.4), 14, WHITE, bold=True)
add_page_number(s2, 2)


# ==========================================
# SLIDE 3: THE SOLUTION
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
add_bg(s3)
add_header(s3, "The Solution", "From meeting notes to meeting execution.", "One transcript in. Real, completed work out — across every tool your team already uses.")

# Left Card
add_card(s3, Inches(0.8), Inches(2.4), Inches(5.0), Inches(3.6), NAVY_SECONDARY)
add_text(s3, "ONE TRANSCRIPT IN", Inches(1.1), Inches(2.7), Inches(4.4), Inches(0.4), 12, CORAL, bold=True)
transcript = "\"...Sarah will update the landing-page mockups. Ravi should send the client recap. Let's document the API decision. And can someone assign the analytics task to Alex... actually, which Alex, marketing or eng?\""
add_text(s3, transcript, Inches(1.1), Inches(3.2), Inches(4.4), Inches(2.5), 22, ICE, 'Cambria', italic=True)

# Right Grid
add_text(s3, "REAL WORK OUT", Inches(6.3), Inches(2.4), Inches(5), Inches(0.4), 12, MINT, bold=True)
grid_w = 2.9
grid_h = 1.3
grid_items = [
    ("T", "Jira ticket created", "Correct project, assignee, priority", MINT, 0, 0),
    ("M", "Gmail draft prepared", "Recipient-ready — never blindly sent", MINT, 1, 0),
    ("D", "Notion page updated", "Decision and context captured", MINT, 0, 1),
    ("?", "Clarification requested", "Ambiguous item paused, not guessed", GOLD, 1, 1)
]

for (icon, title, desc, dot_col, col, row) in grid_items:
    x = 6.3 + (col * (grid_w + 0.3))
    y = 2.8 + (row * (grid_h + 0.3))
    add_card(s3, Inches(x), Inches(y), Inches(grid_w), Inches(grid_h))
    add_circle_icon(s3, Inches(x+0.2), Inches(y+0.2), Inches(0.4), icon)
    # status dot
    dot = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+grid_w-0.3), Inches(y+0.2), Inches(0.15), Inches(0.15))
    dot.fill.solid()
    dot.fill.fore_color.rgb = dot_col
    dot.line.fill.background()
    
    add_text(s3, title, Inches(x+0.7), Inches(y+0.2), Inches(2.0), Inches(0.3), 14, WHITE, bold=True)
    add_text(s3, desc, Inches(x+0.2), Inches(y+0.7), Inches(2.5), Inches(0.5), 12, ICE)

add_card(s3, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.5), NAVY_PRIMARY, MINT)
add_text(s3, "4 action items · 1 ambiguous · 1 completion report", Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.4), 16, MINT, bold=True, align=PP_ALIGN.CENTER)
add_page_number(s3, 3)

# ==========================================
# SLIDE 4: HOW IT WORKS
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
add_bg(s4)
add_header(s4, "How It Works", "An autonomous, confidence-aware execution loop.", "The system decides what to do, acts through real tools, watches what happens, and adapts when context is incomplete.")

step_w = 2.2
step_h = 1.8
steps = [
    ("1", "INGEST", "Meeting transcript"),
    ("2", "EXTRACT", "Structured action items"),
    ("3", "PLAN", "Tool + owner + priority"),
    ("4", "EXECUTE", "Jira · Gmail · Notion"),
    ("5", "REPORT", "Links + final status")
]

for i, (num, title, desc) in enumerate(steps):
    x = 0.8 + (i * (step_w + 0.2))
    y = 3.0
    add_card(s4, Inches(x), Inches(y), Inches(step_w), Inches(step_h))
    add_circle_icon(s4, Inches(x+0.2), Inches(y+0.2), Inches(0.4), num)
    add_text(s4, title, Inches(x+0.2), Inches(y+0.8), Inches(1.8), Inches(0.3), 16, WHITE, bold=True)
    add_text(s4, desc, Inches(x+0.2), Inches(y+1.2), Inches(1.8), Inches(0.5), 13, ICE)
    
    if i < 4:
        # Arrow
        add_text(s4, "→", Inches(x+step_w), Inches(y+0.7), Inches(0.2), Inches(0.4), 24, ICE)

add_card(s4, Inches(0.8), Inches(5.2), Inches(11.8), Inches(0.8), NAVY_SECONDARY, CORAL)
add_text(s4, "Low confidence or unresolved owner? Pause that one item → ask a human in Slack → apply the answer → resume execution.", Inches(1.0), Inches(5.45), Inches(11.4), Inches(0.4), 16, WHITE, bold=True)

add_text(s4, "This loop — decide, act, observe, adapt — is what makes it agentic, not just an LLM response inside a dashboard.", Inches(0.8), Inches(6.5), Inches(11.8), Inches(0.4), 14, ICE, italic=True, align=PP_ALIGN.CENTER)
add_page_number(s4, 4)

# ==========================================
# SLIDE 5: WHY THIS STANDS OUT
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
add_bg(s5)
add_header(s5, "Why This Stands Out", "A meeting assistant that proves agentic behaviour.", "Not better extraction — autonomous execution, uncertainty handling, and memory.")

# Table headers
add_text(s5, "CAPABILITY", Inches(0.8), Inches(2.6), Inches(3.0), Inches(0.4), 14, WHITE, bold=True)
add_text(s5, "TYPICAL NOTES BOT", Inches(4.0), Inches(2.6), Inches(4.0), Inches(0.4), 14, CORAL, bold=True)
add_text(s5, "MEETING TO MOTION", Inches(8.5), Inches(2.6), Inches(4.0), Inches(0.4), 14, MINT, bold=True)

rows = [
    ("Planning", "Outputs a static task list", "Creates the actual task in the tool"),
    ("Tool use", "Ends after summarization", "Runs a multi-step execution workflow"),
    ("Self-correction", "Guesses missing context / hallucinate", "Pauses and asks for human clarification"),
    ("Memory", "No persistent team context", "Applies past team conventions via graph"),
    ("Autonomous loop", "Text-only result", "Auditable links and live execution statuses")
]

y = 3.2
for cap, bot, m2m in rows:
    add_text(s5, cap, Inches(0.8), Inches(y), Inches(3.0), Inches(0.4), 16, WHITE, bold=True)
    
    add_card(s5, Inches(4.0), Inches(y-0.1), Inches(4.0), Inches(0.5), RED_TINT)
    add_text(s5, bot, Inches(4.1), Inches(y), Inches(3.8), Inches(0.4), 14, ICE)
    
    add_card(s5, Inches(8.5), Inches(y-0.1), Inches(4.0), Inches(0.5), NAVY_SECONDARY)
    add_text(s5, m2m, Inches(8.6), Inches(y), Inches(3.8), Inches(0.4), 14, MINT, bold=True)
    y += 0.65

add_text(s5, "EXECUTION IS THE MOAT", Inches(0.8), Inches(6.6), Inches(11.8), Inches(0.4), 20, CORAL, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_page_number(s5, 5)

# ==========================================
# SLIDE 6: TECHNICAL ARCHITECTURE
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
add_bg(s6)
add_header(s6, "Technical Architecture", "Modular enough for four people. Simple enough for 24 hours.", "Every module exchanges the same structured ActionItem contract, so the team builds in parallel — not in sequence.")

layers = [
    ("INPUT", [("Sample transcript", "Free"), ("Optional audio - faster-whisper", "Free, local")]),
    ("REASONING", [("Gemini API", "Free tier"), ("LangGraph orchestration", "Free/OSS"), ("JSON/SQLite team memory", "Free"), ("Groq fallback", "Speed")]),
    ("EXECUTION", [("Jira REST API", "Free Cloud"), ("Gmail API - drafts only", "Free, quota"), ("Notion API", "Fully free"), ("Slack API", "Free workspace")]),
    ("EXPERIENCE", [("React/Vite dashboard", "Live status"), ("Completion report", "Links + outcomes")])
]

y = 2.8
for label, tags in layers:
    add_card(s6, Inches(0.8), Inches(y), Inches(11.7), Inches(0.7), NAVY_SECONDARY)
    add_text(s6, label, Inches(1.0), Inches(y+0.25), Inches(1.5), Inches(0.3), 14, CORAL, bold=True)
    
    x_offset = 2.5
    for text, sub in tags:
        add_text(s6, text, Inches(x_offset), Inches(y+0.15), Inches(2.0), Inches(0.3), 14, WHITE, bold=True)
        add_text(s6, sub, Inches(x_offset), Inches(y+0.4), Inches(2.0), Inches(0.3), 11, ICE)
        x_offset += 2.2
    y += 0.9

add_card(s6, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5), NAVY_PRIMARY, MINT)
add_text(s6, "Safety by design: Gmail creates drafts, never sends · unclear items are paused, never guessed · every action returns a status and a link", Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.3), 12, MINT, align=PP_ALIGN.CENTER)
add_page_number(s6, 6)


# ==========================================
# SLIDE 7: 24-HOUR FEASIBILITY
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
add_bg(s7)
add_header(s7, "24-Hour Feasibility", "A focused MVP with a clear division of work.", "One clean end-to-end path beats a broad but fragile demo.")

# Timeline
tl_y = 2.8
s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(tl_y+0.2), Inches(10.9), Inches(0.05)).fill.solid()
s7.shapes[-1].fill.fore_color.rgb = ICE

timeline = [("Setup (0-1h)", "Contracts + auth"), ("Foundations (1-6h)", "Modules work alone"), ("Core Build (6-14h)", "Real data + clarity"), ("Integrate (14-18h)", "Full pipeline"), ("Polish (18-24h)", "Demo ready")]
for i, (title, sub) in enumerate(timeline):
    x = 1.0 + (i * 2.5)
    add_circle_icon(s7, Inches(x+0.3), Inches(tl_y+0.1), Inches(0.3), "", CORAL)
    add_text(s7, title, Inches(x-0.2), Inches(tl_y+0.5), Inches(1.5), Inches(0.3), 12, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s7, sub, Inches(x-0.2), Inches(tl_y+0.8), Inches(1.5), Inches(0.3), 11, ICE, align=PP_ALIGN.CENTER)

# Team split
add_text(s7, "4-PERSON TEAM SPLIT", Inches(0.8), Inches(4.0), Inches(5), Inches(0.3), 12, CORAL, bold=True)
team = [
    ("01 The Brain", "Extraction/routing/memory"),
    ("02 Docs & Tickets", "Jira/Notion APIs"),
    ("03 Comms & Clarity", "Gmail/Slack loop"),
    ("04 Glue & Demo", "Orchestrator/UI/testing")
]
for i, (t, d) in enumerate(team):
    add_card(s7, Inches(0.8 + i*2.9), Inches(4.4), Inches(2.7), Inches(0.8))
    add_text(s7, t, Inches(1.0 + i*2.9), Inches(4.5), Inches(2.5), Inches(0.3), 14, WHITE, bold=True)
    add_text(s7, d, Inches(1.0 + i*2.9), Inches(4.8), Inches(2.5), Inches(0.3), 12, ICE)

# Bottom Split
add_card(s7, Inches(0.8), Inches(5.6), Inches(5.7), Inches(1.0), NAVY_SECONDARY, MINT)
add_text(s7, "LOCKED MVP", Inches(1.0), Inches(5.8), Inches(2), Inches(0.3), 12, MINT, bold=True)
add_text(s7, "1 Jira ticket · 1 Gmail draft · 1 Notion update · 1 Slack clarification resolved · 1 completion report", Inches(1.0), Inches(6.1), Inches(5.3), Inches(0.4), 12, WHITE)

add_card(s7, Inches(6.8), Inches(5.6), Inches(5.7), Inches(1.0), NAVY_SECONDARY, GOLD)
add_text(s7, "FALLBACK PLAN", Inches(7.0), Inches(5.8), Inches(2), Inches(0.3), 12, GOLD, bold=True)
add_text(s7, "Recorded demo + cached extraction results protect the pitch from rate limits or network drops.", Inches(7.0), Inches(6.1), Inches(5.3), Inches(0.4), 12, WHITE)
add_page_number(s7, 7)


# ==========================================
# SLIDE 8: PROOF OF WORK
# ==========================================
s8 = prs.slides.add_slide(blank_layout)
add_bg(s8)
add_header(s8, "Proof Of Work", "Not a concept. A running prototype.", "Real dashboard runs, real Slack clarifications, real commit history — built during this hackathon.")

p_w = 3.7
p_h = 2.2
panels = [
    ("Live Pipeline Dashboard", "Streamlit/React UI running the real extract → execute loop on a pasted transcript."),
    ("Real Slack Clarification", "The agent pausing on an ambiguous owner and asking a human in #test-channel — live, not simulated."),
    ("Active Commit History", "Multiple team members shipping real commits — extraction, LLM fixes, schema adherence.")
]

for i, (title, desc) in enumerate(panels):
    x = 0.8 + (i * (p_w + 0.3))
    # Placeholder for image
    add_card(s8, Inches(x), Inches(2.6), Inches(p_w), Inches(p_h), NAVY_SECONDARY, ICE)
    add_text(s8, "[Insert Screenshot Here]", Inches(x), Inches(3.5), Inches(p_w), Inches(0.4), 14, ICE, align=PP_ALIGN.CENTER)
    
    add_text(s8, title, Inches(x), Inches(5.0), Inches(p_w), Inches(0.3), 16, WHITE, bold=True)
    add_text(s8, desc, Inches(x), Inches(5.4), Inches(p_w), Inches(0.8), 13, ICE)

add_card(s8, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5), NAVY_PRIMARY, MINT)
add_text(s8, "STATUS: extraction working · Slack clarification loop live · dashboard rendering real results · team shipping in parallel", Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.3), 13, MINT, bold=True, align=PP_ALIGN.CENTER)
add_page_number(s8, 8)


# ==========================================
# SLIDE 9: IMPACT & CLOSING
# ==========================================
s9 = prs.slides.add_slide(blank_layout)
add_bg(s9)
add_header(s9, "Impact & Closing", "Every meeting ends with visible, completed movement.", "MeetingToMotion shrinks the distance between 'we decided' and 'it is done.'")

# Stats
add_card(s9, Inches(0.8), Inches(2.5), Inches(2.7), Inches(1.5))
add_text(s9, "4", Inches(0.8), Inches(2.7), Inches(2.7), Inches(0.8), 64, GOLD, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_text(s9, "Action items processed", Inches(0.8), Inches(3.6), Inches(2.7), Inches(0.3), 14, ICE, align=PP_ALIGN.CENTER)

add_card(s9, Inches(3.8), Inches(2.5), Inches(2.7), Inches(1.5))
add_text(s9, "3", Inches(3.8), Inches(2.7), Inches(2.7), Inches(0.8), 64, GOLD, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_text(s9, "Work tools updated", Inches(3.8), Inches(3.6), Inches(2.7), Inches(0.3), 14, ICE, align=PP_ALIGN.CENTER)

add_card(s9, Inches(6.8), Inches(2.5), Inches(2.7), Inches(1.5))
add_text(s9, "1", Inches(6.8), Inches(2.7), Inches(2.7), Inches(0.8), 64, GOLD, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_text(s9, "Ambiguity safely resolved", Inches(6.8), Inches(3.6), Inches(2.7), Inches(0.3), 14, ICE, align=PP_ALIGN.CENTER)

add_card(s9, Inches(9.8), Inches(2.5), Inches(2.7), Inches(1.5))
add_text(s9, "0", Inches(9.8), Inches(2.7), Inches(2.7), Inches(0.8), 64, MINT, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_text(s9, "Manual copy-paste steps", Inches(9.8), Inches(3.6), Inches(2.7), Inches(0.3), 14, ICE, align=PP_ALIGN.CENTER)

# Bottom section
add_text(s9, "WHY IT MATTERS", Inches(0.8), Inches(4.5), Inches(4), Inches(0.4), 14, CORAL, bold=True)
reasons = [
    ("Faster execution", "Decisions immediately become trackable work"),
    ("Fewer dropped actions", "Every item receives an outcome or a clarification status"),
    ("Less context switching", "Teams stop manually translating notes across tools")
]
for i, (title, desc) in enumerate(reasons):
    add_circle_icon(s9, Inches(0.8), Inches(5.0 + i*0.6), Inches(0.3), "")
    add_text(s9, title, Inches(1.2), Inches(4.9 + i*0.6), Inches(3.0), Inches(0.3), 14, WHITE, bold=True)
    add_text(s9, desc, Inches(1.2), Inches(5.2 + i*0.6), Inches(4.0), Inches(0.3), 12, ICE)

add_card(s9, Inches(5.8), Inches(4.6), Inches(6.7), Inches(1.5), NAVY_SECONDARY, MINT)
add_text(s9, "THE DEMO MOMENT", Inches(6.0), Inches(4.8), Inches(5), Inches(0.3), 12, MINT, bold=True)
add_text(s9, "Paste transcript → watch the agent create real work → answer one Slack clarification → open the final completion report.", Inches(6.0), Inches(5.2), Inches(6.3), Inches(0.6), 16, WHITE, italic=True)
add_text(s9, "LIVE · AUDITABLE · AGENTIC", Inches(6.0), Inches(5.7), Inches(5), Inches(0.3), 11, ICE, bold=True)

add_text(s9, "M→M   We built the execution. Extraction is commoditized.", Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4), 16, WHITE, 'Cambria', bold=True, align=PP_ALIGN.CENTER)
add_page_number(s9, 9)

prs.save("MeetingToMotion_PitchDeck.pptx")
print("Pitch deck successfully generated: MeetingToMotion_PitchDeck.pptx")
